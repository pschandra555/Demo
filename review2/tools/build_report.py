"""Builds the Review II report (DOCX) and the screenshot checklist (DOCX).

Usage:  python3 tools/build_report.py
Screenshots: save them as screenshots/fig_02.png, fig_03.png ... (figure numbers from the
checklist) and run the script again; every placeholder that has a matching file is replaced
by the real screenshot.
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Emu

sys.path.insert(0, os.path.dirname(__file__))
from figures import FIGURES, BY_KEY, fig  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SQL = os.path.join(ROOT, 'sql')
SHOTS = os.path.join(ROOT, 'screenshots')

STUDENT = ('P. SARATH CHANDRA', '24BYB0047')   # header of every page - change if needed
TEAM = [('P. Sarath Chandra', '24BYB0047'), ('Palli Pradeep', '24BYB0025'), ('M. Prasanna', '24BDE0078')]
TITLE = 'Vehicle Rental Management System'

GREEN = RGBColor(0x1E, 0x6B, 0x52)
DARK = RGBColor(0x26, 0x28, 0x2E)
GREY = RGBColor(0x6B, 0x6E, 0x75)
BODY_FONT = 'Calibri'
MONO = 'Consolas'


# ------------------------------------------------------------------ helpers
def chunks(fname):
    res, key, buf = {}, None, []
    for line in open(os.path.join(SQL, fname)):
        m = re.match(r'--## (\S+)', line)
        if m:
            if key:
                res[key] = ''.join(buf).strip('\n')
            key, buf = m.group(1), []
        elif key:
            buf.append(line)
    if key:
        res[key] = ''.join(buf).strip('\n')
    return res


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def cell_borders(cell, **kw):
    """kw: top/bottom/left/right = (size_eighths, color, style)"""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge in kw:
            sz, color, style = kw[edge]
            el = OxmlElement(f'w:{edge}')
            el.set(qn('w:val'), style)
            el.set(qn('w:sz'), str(sz))
            el.set(qn('w:space'), '0')
            el.set(qn('w:color'), color)
            borders.append(el)
    tcPr.append(borders)


def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for k, v in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        el = OxmlElement(f'w:{k}')
        el.set(qn('w:w'), str(v))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tcPr.append(mar)


def table_borders(table, color='D0D3D8', sz=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    trPr.append(el)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    trPr.append(el)


def set_col_widths(table, widths_cm):
    table.autofit = False
    tbl = table._tbl
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths_cm[i] * 567)))
    for row in table.rows:
        for i, c in enumerate(row.cells):
            c.width = Cm(widths_cm[i])


def para_fmt(p, before=0, after=4, line=1.1, keep=False, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep
    if align is not None:
        p.alignment = align
    return p


def add_run(p, text, bold=False, italic=False, size=None, color=None, font=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    if font:
        r.font.name = font
        r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    return r


def rich(p, text, size=None, color=None):
    """**bold** and `code` markup inside a paragraph."""
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**'):
            add_run(p, part[2:-2], bold=True, size=size, color=color)
        elif part.startswith('`'):
            add_run(p, part[1:-1], size=(size or 10.5) - 1, font=MONO, color=RGBColor(0x8A, 0x2B, 0x1E))
        else:
            add_run(p, part, size=size, color=color)


class Report:
    def __init__(self):
        self.doc = Document()
        self.placed = []
        st = self.doc.styles['Normal']
        st.font.name = BODY_FONT
        st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)
        st.font.size = Pt(10.5)
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.6)
        sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)
        self._header_footer(sec)

    # ---- page furniture
    def _header_footer(self, sec):
        hp = sec.header.paragraphs[0]
        
        add_run(hp, f'NAME : {STUDENT[0]}', bold=True, size=9.5, color=DARK)
        add_run(hp, f'\t\tREG NO : {STUDENT[1]}', bold=True, size=9.5, color=DARK)
        fp = sec.footer.paragraphs[0]
        fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
        add_run(fp, f'{TITLE}  |  BCSE302L Database Systems  |  Review II', size=8.5, color=GREY)
        add_run(fp, '\tPage ', size=8.5, color=GREY)
        r = add_run(fp, '', size=8.5, color=GREY)
        for kind, text in (('begin', None), (None, 'PAGE'), ('end', None)):
            if kind:
                el = OxmlElement('w:fldChar')
                el.set(qn('w:fldCharType'), kind)
            else:
                el = OxmlElement('w:instrText')
                el.set(qn('xml:space'), 'preserve')
                el.text = text
            r._element.append(el)
        pPr = fp._p.get_or_add_pPr()
        bdr = OxmlElement('w:pBdr')
        top = OxmlElement('w:top')
        for k, v in (('val', 'single'), ('sz', '4'), ('space', '4'), ('color', 'D0D3D8')):
            top.set(qn(f'w:{k}'), v)
        bdr.append(top)
        pPr.append(bdr)

    # ---- text blocks
    def h1(self, number, text, page_break=False):
        p = self.doc.add_paragraph()
        if page_break:
            p.paragraph_format.page_break_before = True
        para_fmt(p, before=4, after=6, keep=True)
        r = add_run(p, f' {number} ', bold=True, size=17, color=DARK)
        rPr = r._element.get_or_add_rPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:fill'), 'F2C230')
        rPr.append(shd)
        add_run(p, f'  {text}', bold=True, size=17, color=DARK)
        return p

    def h2(self, text):
        p = self.doc.add_paragraph()
        para_fmt(p, before=8, after=3, keep=True)
        add_run(p, text, bold=True, size=12.5, color=GREEN)
        return p

    def h3(self, text, sub=None):
        p = self.doc.add_paragraph()
        para_fmt(p, before=7, after=1, keep=True)
        add_run(p, text, bold=True, size=11, color=DARK)
        if sub:
            p2 = self.doc.add_paragraph()
            para_fmt(p2, after=3, keep=True)
            rich(p2, sub, size=9, color=GREY)
        return p

    def para(self, text, size=10.5, after=4, keep=False, align=None):
        p = self.doc.add_paragraph()
        para_fmt(p, after=after, keep=keep, align=align)
        rich(p, text, size=size)
        return p

    def bullets(self, items, size=10.5):
        for it in items:
            p = self.doc.add_paragraph()
            para_fmt(p, after=2)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            add_run(p, '\u25aa  ', color=GREEN, size=size)
            rich(p, it, size=size)

    def table(self, header, rows, widths, size=9, first_col_bold=True, first_col_color=GREEN):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(t)
        hdr = t.rows[0]
        repeat_header(hdr)
        for i, h in enumerate(header):
            c = hdr.cells[i]
            shade(c, '2B2D33')
            cell_margins(c)
            p = c.paragraphs[0]
            para_fmt(p, after=0)
            add_run(p, h, bold=True, size=size, color=RGBColor(0xFF, 0xFF, 0xFF))
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            no_split(t.rows[-1])
            for i, val in enumerate(row):
                c = cells[i]
                cell_margins(c)
                if ri % 2 == 1:
                    shade(c, 'F4F5F6')
                p = c.paragraphs[0]
                para_fmt(p, after=0, line=1.05)
                if i == 0 and first_col_bold:
                    add_run(p, str(val), bold=True, size=size, color=first_col_color)
                else:
                    rich(p, str(val), size=size)
        set_col_widths(t, widths)
        sp = self.doc.add_paragraph()
        para_fmt(sp, after=2)
        sp.paragraph_format.line_spacing = 0.6
        return t

    def code(self, text, size=7):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        shade(c, 'F5F6F3')
        cell_borders(c, left=(18, '1E6B52', 'single'))
        cell_margins(c, top=70, bottom=70, left=140, right=80)
        p = c.paragraphs[0]
        para_fmt(p, after=0, line=1.0)
        lines = text.rstrip('\n').split('\n')
        for i, line in enumerate(lines):
            r = add_run(p, line, size=size, font=MONO, color=RGBColor(0x22, 0x24, 0x28))
            if i < len(lines) - 1:
                r.add_break()
        set_col_widths(t, [17.0])
        sp = self.doc.add_paragraph()
        para_fmt(sp, after=2)
        sp.paragraph_format.line_spacing = 0.6

    def figure(self, key, note=None, width_cm=16.0):
        f = BY_KEY[key]
        expected_no = len(self.placed) + 1
        assert f['no'] == expected_no, f'figure {key} placed out of order ({f["no"]} vs {expected_no})'
        self.placed.append(key)
        shot = os.path.join(SHOTS, f"fig_{f['no']:02d}.png")
        image = f.get('image') and os.path.join(ROOT, f['image'])
        path = image or (shot if os.path.exists(shot) else None)
        if path:
            from PIL import Image
            w, h = Image.open(path).size
            width = Cm(width_cm) if image else Cm(min(16.5, w * 0.0118))
            max_h = Cm(f.get('max_h', 13.5))
            if width * h / w > max_h:
                width = Emu(int(max_h * w / h))
            p = self.doc.add_paragraph()
            para_fmt(p, before=2, after=1, keep=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.add_run().add_picture(path, width=width)
        else:
            t = self.doc.add_table(rows=1, cols=1)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            row = t.rows[0]
            row.height = Cm(round(f['h'] * 0.88, 2))
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            no_split(row)
            c = row.cells[0]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(c, 'FAFAF7')
            dash = (8, 'A7ABB2', 'dashed')
            cell_borders(c, top=dash, bottom=dash, left=dash, right=dash)
            p = c.paragraphs[0]
            para_fmt(p, after=2, keep=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_run(p, f"SCREENSHOT PLACEHOLDER \u2013 Figure {f['no']}", bold=True, size=10, color=GREY)
            p2 = c.add_paragraph()
            para_fmt(p2, after=0, keep=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            src = f.get('src')
            how = f"run block {src[1]} of {src[0]}" if src else f.get('run', '').replace('\n', ' then ')
            add_run(p2, f'Delete this box and paste the screenshot ({how}). Listed as Figure {f["no"]} in the screenshot checklist.',
                    size=8.5, color=GREY)
            set_col_widths(t, [16.0])
        cp = self.doc.add_paragraph()
        para_fmt(cp, before=2, after=6 if not note else 2, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_run(cp, f"Figure {f['no']}: ", bold=True, italic=True, size=9, color=DARK)
        add_run(cp, f['caption'], italic=True, size=9, color=GREY)
        if note:
            np_ = self.doc.add_paragraph()
            para_fmt(np_, after=6)
            rich(np_, note, size=10)

    def page_break(self):
        p = self.doc.add_paragraph()
        para_fmt(p, after=0)
        p.add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        assert len(self.placed) == len(FIGURES), f'only {len(self.placed)} of {len(FIGURES)} figures placed'
        self.doc.save(path)


# ------------------------------------------------------------------ content
def build_report(out_path):
    R = Report()
    d = R.doc
    ddl = chunks('01_ddl.sql')
    ins = chunks('02_insert_data.sql')
    dml = chunks('06_dml_demo.sql')
    qs = chunks('05_queries.sql')
    fn = chunks('07_functions.sql')
    sp = chunks('08_procedures.sql')
    cu = chunks('09_cursors.sql')
    tg = chunks('10_triggers.sql')

    # ---------------- cover
    p = d.add_paragraph()
    para_fmt(p, before=10, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, 'Vellore Institute of Technology', bold=True, size=13, color=DARK)
    p = d.add_paragraph()
    para_fmt(p, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, TITLE, bold=True, size=24, color=GREEN)
    p = d.add_paragraph()
    para_fmt(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, 'Course Project \u2013 Review II Report', bold=True, size=13, color=DARK)
    p = d.add_paragraph()
    para_fmt(p, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, '\u2500' * 30, color=RGBColor(0xF2, 0xC2, 0x30), size=11)

    t = R.table(['S.No', 'Team member', 'Register number'],
                [[str(i + 1), f'**{n}**', r] for i, (n, r) in enumerate(TEAM)], [1.6, 7.4, 4.8],
                size=10, first_col_bold=False)
    for row in t.rows[1:]:
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    R.table(['Item', 'Details'],
            [['Course', 'BCSE302L \u2013 Database Systems'],
             ['Milestone', 'Review II \u2013 Database Implementation'],
             ['Slot', 'L1 + L2'],
             ['Faculty', 'Dr. Ram Kumar T'],
             ['Database', 'Oracle AI Database 26ai Free (release 23.26), SQL*Plus']],
            [4.0, 9.8], size=10, first_col_color=DARK)

    R.h2('Contents')
    R.table(['Section', 'Title', 'Rubric marks'],
            [['0', 'Implementation summary and link to Review I', '\u2013'],
             ['1', 'DDL and constraints', '2'],
             ['2', 'DML and data population', '1'],
             ['3', 'SQL queries', '2'],
             ['4', 'Procedures and functions', '2'],
             ['5', 'Cursors and triggers', '2'],
             ['6', 'Real-world transaction flow and business rules', '\u2013'],
             ['7', 'Viva preparation: SQL / PL/SQL understanding', '1'],
             ['8', 'Conclusion', '\u2013']],
            [2.0, 11.5, 3.5], size=9.5)

    # ---------------- section 0
    R.h1('0', 'Implementation Summary', page_break=True)
    R.para('In Review I we designed the database for a multi-branch vehicle rental business: eight entities, '
           'ten one-to-many relationships, twelve functional requirements (FR-01 to FR-12) and eight business '
           'rules (BR-01 to BR-08). In this review the same design is implemented in Oracle. Table names, column '
           'names, data types and constraints are exactly those listed in Section 4.3 of the Review I report; no '
           'table or column was added or removed. The business rules are enforced inside the database through '
           'constraints, functions, procedures and triggers.')
    R.table(['Rubric component', 'Section', 'What is implemented'],
            [['DDL & Constraints (2)', '1', '8 tables, 4 sequences; 8 PRIMARY KEY, 10 FOREIGN KEY, 10 UNIQUE, 23 CHECK, '
              '34 NOT NULL constraints and 14 DEFAULT values. 8 constraint tests.'],
             ['DML & Data Population (1)', '2', '140 rows (10 to 45 per table). INSERT with DEFAULT, UPDATE, and '
              'DELETE with ON DELETE CASCADE.'],
             ['SQL Queries (2)', '3', '16 queries: inner/outer/self joins, GROUP BY, HAVING, nested and correlated '
              'subqueries, MINUS, UNION ALL, CASE, date functions, top-N.'],
             ['Procedures & Functions (2)', '4', '6 functions (age, rental days, rent, late fee, availability, balance) '
              'and 5 procedures (register, book, pickup, return, cancel).'],
             ['Cursor & Triggers (2)', '5', '2 cursor procedures (explicit cursor and cursor FOR loop) and 4 row-level '
              'triggers.'],
             ['Viva / Understanding (1)', '6, 7', 'Transaction flow table, business-rule matrix and likely viva questions.']],
            [4.4, 1.6, 11.0])
    R.h2('Schema from Review I')
    R.figure('er', width_cm=8.5)
    R.code('BRANCH (branch_id, branch_name, city, address, phone, email)\n'
           'VEHICLE_CATEGORY (category_id, category_name, daily_rate, late_fee_per_hr, seating_capacity, security_deposit)\n'
           'VEHICLE (vehicle_id, reg_number, brand, model, mfg_year, fuel_type, transmission, odometer, status,\n'
           '         category_id -> VEHICLE_CATEGORY, branch_id -> BRANCH)\n'
           'CUSTOMER (customer_id, name, phone, email, license_no, license_expiry, dob, address, registered_on)\n'
           'EMPLOYEE (emp_id, name, role, phone, email, salary, hire_date, branch_id -> BRANCH)\n'
           'BOOKING (booking_id, customer_id -> CUSTOMER, vehicle_id -> VEHICLE, emp_id -> EMPLOYEE,\n'
           '         pickup_branch_id -> BRANCH, return_branch_id -> BRANCH, booking_date, start_date, end_date,\n'
           '         actual_return_date, start_odometer, end_odometer, status)\n'
           'PAYMENT (payment_id, booking_id -> BOOKING, amount, payment_date, method, payment_type, txn_status)\n'
           'MAINTENANCE (maint_id, vehicle_id -> VEHICLE, service_date, service_type, description, cost, maint_status)',
           size=7.5)
    R.h2('Order of execution')
    R.table(['Step', 'Script (folder sql/)', 'What it does', 'Figures'],
            [['0', '00_setup_display.sql', 'Display settings (line size, column widths, date format). Run once per session.', '\u2013'],
             ['1', '01_ddl.sql', 'Creates the 4 sequences and 8 tables (parent tables first).', fig('ddl_run')[7:]],
             ['2', '02_insert_data.sql', 'Loads the 140 sample rows and sets current vehicle states; COMMIT.', '\u2013'],
             ['3', '03_check_schema_data.sql', 'DESC, constraint summary, row counts and table listings (read only).',
              f"{BY_KEY['sc_desc']['no']}\u2013{BY_KEY['sc_fk']['no']}, {BY_KEY['data_counts']['no']}\u2013{BY_KEY['data_booking']['no']}"],
             ['4', '04_constraint_tests.sql', 'Eight statements that must fail; ends with ROLLBACK.',
              f"{BY_KEY['ct1']['no']}\u2013{BY_KEY['ct2']['no']}"],
             ['5', '05_queries.sql', 'Queries Q1\u2013Q16 on the loaded data.', f"{BY_KEY['q1']['no']}\u2013{BY_KEY['q16']['no']}"],
             ['6', '06_dml_demo.sql', 'INSERT, UPDATE and DELETE demonstration (changes a few rows).',
              f"{BY_KEY['dml_insert']['no']}\u2013{BY_KEY['dml_delete']['no']}"],
             ['7', '07 to 10 (.sql)', 'Functions, procedures, cursors and triggers (CREATE OR REPLACE).', '\u2013'],
             ['8', '11_transactions.sql', 'Runs the PL/SQL objects block by block, in order.',
              f"{BY_KEY['objects']['no']}\u2013{BY_KEY['rules']['no']}"]],
            [1.2, 4.3, 9.3, 2.2], size=8.5)
    R.para('**About the screenshots.** All 42 screenshots in this report come from one run of these scripts, in this order, '
           'on a fresh schema in Oracle AI Database 26ai Free using SQL*Plus (13-Oct-2026). Each image shows the commands as '
           'echoed by SQL*Plus (SET ECHO ON) followed by the unedited output. Outputs that use SYSDATE (Q16, the overdue '
           'report and the demo booking dates) depend on the day the scripts are run.', size=9.5)
    R.para('99_cleanup.sql drops every object so the whole sequence can be repeated on fresh data. '
           'Steps 5 and 6 must stay in this order because the DML demonstration changes rows that the queries read.',
           size=9.5)

    # ---------------- section 1
    R.h1('1', 'DDL and Constraints')
    R.para('The tables are created parent first (BRANCH and VEHICLE_CATEGORY, then VEHICLE, CUSTOMER and EMPLOYEE, '
           'then BOOKING, and finally PAYMENT and MAINTENANCE) so that every foreign key points to a table that '
           'already exists. Every constraint is given a name with a prefix (`pk_`, `fk_`, `uq_`, `ck_`). Because of '
           'this, an error message such as ORA-02290 directly shows which rule was broken. The four sequences are used '
           'only by the procedures; they start after the IDs used in the sample data.')
    R.h2('1.1 Sequences and table definitions')
    R.code(ddl['DDL_SEQ'])
    R.code(ddl['DDL_MASTER'])
    R.code(ddl['DDL_PEOPLE'])
    R.code(ddl['DDL_BOOKING'])
    R.code(ddl['DDL_CHILD'])
    R.figure('ddl_run')
    R.figure('sc_desc', note='DESC confirms that the columns, data types and NOT NULL settings match the Review I '
                             'design. The primary key column is also shown as NOT NULL.')
    R.h2('1.2 Constraints used')
    R.table(['Constraint', 'Count', 'Where it is used', 'Purpose'],
            [['PRIMARY KEY', '8', 'One per table: branch_id, category_id, vehicle_id, customer_id, emp_id, booking_id, '
              'payment_id, maint_id', 'Identifies every row uniquely.'],
             ['FOREIGN KEY', '10', 'vehicle (2), employee (1), booking (5), payment (1), maintenance (1)',
              'One per relationship in the ER diagram. PAYMENT uses ON DELETE CASCADE, as planned in Review I.'],
             ['UNIQUE', '10', 'branch name/phone/email, category name, reg_number, customer phone/email/licence, '
              'employee phone/email', 'Stops duplicate vehicles, customers and contact details.'],
             ['CHECK', '23', 'Allowed values of status, fuel, role, method, payment type; rate > 0; seats 1\u201315; '
              'mfg_year >= 2010; phone length 10; end_date > start_date; end_odometer >= start_odometer (BR-08)',
              'Keeps values inside the business limits.'],
             ['NOT NULL', '34', 'Names, phones, dates and every foreign key on the "many" side', 'Total participation '
              'from the ER diagram: e.g. a vehicle cannot exist without a category and a branch.'],
             ['DEFAULT', '14', "status 'AVAILABLE'/'BOOKED'/'SCHEDULED', transmission 'MANUAL', odometer 0, "
              "late fee 100, deposit 2000, txn_status 'SUCCESS', cost 0, dates SYSDATE", 'Sensible value when a column is left out.']],
            [2.4, 1.3, 7.5, 5.8], size=8.5)
    R.para('Age \u2265 21 and licence validity depend on the current date. Oracle does not allow SYSDATE inside a CHECK '
           'constraint, so (as noted in Review I) these two rules are enforced by triggers in Section 5.', size=10)
    R.figure('sc_count')
    R.figure('sc_fk')
    R.h2('1.3 Constraint testing (each statement must fail)')
    R.table(['#', 'Statement (04_constraint_tests.sql)', 'Constraint', 'Expected Oracle error'],
            [['1', 'Insert vehicle with reg_number TN09AB1001 (already used)', 'UNIQUE', 'ORA-00001 UQ_VEHICLE_REG'],
             ['2', "Insert vehicle with fuel_type 'HYBRID'", 'CHECK', 'ORA-02290 CK_VEHICLE_FUEL'],
             ['3', "Insert branch with phone '98401' (5 digits)", 'CHECK', 'ORA-02290 CK_BRANCH_PHONE'],
             ['4', 'Insert customer without license_no', 'NOT NULL', 'ORA-01400 CUSTOMER.LICENSE_NO'],
             ['5', 'Insert payment for booking 999', 'FOREIGN KEY (child)', 'ORA-02291 FK_PAYMENT_BOOKING'],
             ['6', 'Delete customer 1, who has bookings', 'FOREIGN KEY (parent)', 'ORA-02292 FK_BOOKING_CUSTOMER'],
             ['7', 'Set end_date one day before start_date', 'Table CHECK', 'ORA-02290 CK_BOOKING_DATES'],
             ['8', 'Set end_odometer below start_odometer (BR-08)', 'Table CHECK', 'ORA-02290 CK_BOOKING_ODOMETER']],
            [0.8, 7.6, 3.4, 5.2], size=8.5)
    R.para('Each test statement runs inside a small PL/SQL block (BEGIN ... EXCEPTION WHEN OTHERS ... END) that catches '
           'the error and prints the first line of the error message. A line starting with **Rejected:** followed by the Oracle error means the '
           'constraint worked and the bad data was refused. If a constraint were missing, the block would print '
           '"Accepted - constraint NOT working" instead. The trigger and business-rule tests in Sections 5 and 6 use the same pattern.', size=10)
    R.figure('ct1')
    R.figure('ct2')

    # ---------------- section 2
    R.h1('2', 'DML and Data Population')
    R.h2('2.1 Sample data (02_insert_data.sql)')
    R.para('The sample data describes a fictional chain, "DriveEase", with ten branches in six cities. It was written so '
           'that the data itself follows the business rules, not just the constraints:')
    R.bullets(['Each booking starts at the branch where the vehicle was parked at that time, and is handled by the agent of that branch.',
               'Odometer readings follow on from one rental to the next; a vehicle\'s current odometer equals the end reading of its last rental.',
               'Every FINAL payment equals rent minus advance, and every LATE_FEE equals hours late \u00d7 the category rate (BR-04, BR-05).',
               'The two cancellations follow BR-06: booking 21 was cancelled 2.7 days early (full refund), booking 22 five hours before pickup (50%).',
               'The data covers all states: 16 completed, 2 ongoing, 2 booked and 2 cancelled bookings; vehicles AVAILABLE, RENTED, '
               'MAINTENANCE and RETIRED; payments SUCCESS, PENDING and FAILED; three one-way rentals (BR-07).'],
              size=10)
    R.table(['Table', 'Rows', 'Table', 'Rows'],
            [['BRANCH', '10', 'EMPLOYEE', '15'], ['VEHICLE_CATEGORY', '10', 'BOOKING', '22'],
             ['VEHICLE', '14', 'PAYMENT', '45'], ['CUSTOMER', '12', 'MAINTENANCE', '12']],
            [4.5, 2.2, 4.5, 2.2], size=9)
    R.para('Representative INSERT statements are shown below; the full script has one INSERT per row.', size=10, keep=True)
    bk = ins['INS_BOOKING'].split('\n')
    pay = [l for l in ins['INS_PAYMENT'].split('\n') if re.match(r'INSERT INTO payment VALUES \(\d+, 2,', l)]
    R.code('\n'.join(ins['INS_BRANCH'].split('\n')[:2]) + '\n\n' + ins['INS_CATEGORY'] + '\n\n' +
           '\n'.join(ins['INS_VEHICLE'].split('\n')[2:4]) + '\n\n' +
           '\n'.join(ins['INS_CUSTOMER'].split('\n')[:2]) + '\n\n' +
           '\n'.join(bk[:2] + bk[5:8]) + '\n\n' + '-- payments of booking 2 (advance, final, late fee)\n' + '\n'.join(pay) +
           '\n\n' + ins['INS_STATUS'], size=7)
    R.figure('data_counts')
    R.figure('data_vehicle')
    R.figure('data_booking')
    R.h2('2.2 INSERT, UPDATE and DELETE operations (06_dml_demo.sql)')
    R.para('These statements are run after the queries of Section 3 because they change a few rows. Each screenshot shows the statements (with their comments) and the result.')
    R.figure('dml_insert')
    R.figure('dml_update')
    R.figure('dml_delete')
    R.table(['Operation', 'Table', 'Effect'],
            [['INSERT', 'vehicle', 'Vehicle 15 (Kia Seltos) added at branch 2 with default transmission, odometer and status.'],
             ['UPDATE', 'employee', 'Salaries of agents 2 and 5 raised by 10%.'],
             ['UPDATE', 'customer', 'Customer 7\'s phone and address changed (UNIQUE on phone still holds).'],
             ['DELETE', 'booking \u2192 payment', 'Cancelled booking 22 removed; its 2 payments are removed automatically by ON DELETE CASCADE.']],
            [2.6, 3.4, 11.0], size=9)

    # ---------------- section 3
    R.h1('3', 'SQL Queries')
    R.para('The queries answer questions that branch agents and the head office actually ask (FR-03, FR-11, FR-12 of '
           'Review I). Each screenshot shows the query as echoed by SQL*Plus together with its output. Dates are displayed in '
           'the DD-MON-YY HH24:MI format set by 00_setup_display.sql.')
    qmeta = {
        'Q1': ('Which vehicles can be rented in Chennai from 21-Oct 10:00 to 23-Oct 10:00? (FR-03)',
               '3-table JOIN, correlated NOT EXISTS subquery, date range overlap | **Tables:** vehicle, vehicle_category, branch, booking',
               'Vehicle 1 (Swift) is missing from the result because booking 19 overlaps these dates, and the retired Indica is filtered out by status.'),
        'Q2': ('Fleet status of every branch (FR-11, fleet utilisation)',
               'LEFT JOIN, COUNT, conditional aggregation with SUM(CASE) | **Tables:** branch, vehicle',
               'The LEFT JOIN keeps the Hyderabad Hitech City branch, which has no vehicle at the moment because its Fortuner was returned at Secunderabad.'),
        'Q3': ('Net revenue collected by each pickup branch (FR-11)',
               '3-table JOIN, SUM with CASE (refunds subtracted), COUNT DISTINCT | **Tables:** payment, booking, branch', None),
        'Q4': ('Month-wise collection by payment type (FR-11)',
               'TO_CHAR date function, GROUP BY expression, conditional SUM | **Tables:** payment', None),
        'Q5': ('Top three customers by amount paid (FR-11)',
               'JOIN, GROUP BY, ORDER BY inside an inline view, ROWNUM for top-N | **Tables:** customer, booking, payment', None),
        'Q6': ('Customers who paid more than the average customer',
               'Nested subquery in HAVING (average of per-customer totals) | **Tables:** customer, booking, payment', None),
        'Q7': ('Registered customers who never made a booking',
               'Set operation MINUS | **Tables:** customer, booking', None),
        'Q8': ('Rentals, average rental length and average distance per category',
               'Two LEFT JOINs, COUNT DISTINCT, AVG of date difference and odometer difference | **Tables:** vehicle_category, vehicle, booking', None),
        'Q9': ('One-way rentals with pickup and return branch names (BR-07)',
               'Self-join: BRANCH joined twice with different aliases | **Tables:** booking, customer, vehicle, branch', None),
        'Q10': ('Late returns, hours late and late fee billed (BR-05)',
                'Date arithmetic (minutes \u2192 started hours), LEFT JOIN with extra join condition, GROUP BY | **Tables:** booking, customer, vehicle, vehicle_category, payment',
                'HOURS_LATE \u00d7 LATE_FEE_PER_HR equals FEE_BILLED for every row, which shows that the stored late-fee payments follow BR-05.'),
        'Q11': ('Payments that are pending or failed (pending dues, FR-11)',
                '4-table JOIN, IN list | **Tables:** payment, booking, customer, vehicle', None),
        'Q12': ('Cancelled bookings with advance and refund (BR-06)',
                'LEFT JOIN, MAX and SUM with CASE | **Tables:** booking, customer, payment',
                'Booking 21 was cancelled more than 24 hours before pickup and got a full refund; booking 22 was cancelled five hours before and got 50%.'),
        'Q13': ('Vehicles whose maintenance cost so far is above Rs 5000 (FR-09)',
                'LEFT JOIN, NVL, string concatenation, HAVING | **Tables:** vehicle, maintenance', None),
        'Q14': ('Vehicles that are out of service right now, with the reason',
                'UNION ALL of two joins, string concatenation, TO_CHAR | **Tables:** vehicle, booking, customer, maintenance', None),
        'Q15': ('Agents who handled two or more bookings',
                'JOIN and LEFT JOIN, WHERE before grouping, HAVING after grouping | **Tables:** employee, branch, booking', None),
        'Q16': ('Customers whose licence expires within the next 12 months (BR-01 follow-up)',
                'SYSDATE, ADD_MONTHS, BETWEEN, TRUNC of a date difference | **Tables:** customer',
                'These customers can still rent, but trg_booking_validate will refuse any booking that ends after their expiry date (Section 5).'),
    }
    for i in range(1, 17):
        k = f'Q{i}'
        title, concepts, note = qmeta[k]
        R.h3(f'{k}. {title}', sub=f'**Concepts:** {concepts}')
        R.figure(f'q{i}', note=note)

    # ---------------- section 4
    R.h1('4', 'Procedures and Functions')
    R.para('Functions return one value and are used inside SQL statements, procedures and triggers. Procedures carry '
           'out a complete business operation: they check the rules, change the data and end with COMMIT. A rule '
           'violation is reported with RAISE_APPLICATION_ERROR, so nothing is saved and the user sees a clear message. '
           'All objects are created with CREATE OR REPLACE and compiled without errors.')
    R.figure('objects')
    R.h2('4.1 Functions (07_functions.sql)')
    R.table(['Function', 'Returns', 'Used by'],
            [['fn_age', 'Age in completed years on a date (derived attribute of CUSTOMER)', 'both validation triggers'],
             ['fn_rental_days', 'Rental days, part of a day = full day, minimum 1 (derived attribute of BOOKING)', 'fn_calculate_rent'],
             ['fn_calculate_rent', 'Rental days \u00d7 daily rate (BR-04)', 'sp_create_booking, sp_return_vehicle'],
             ['fn_late_fee', 'Started hours after end_date \u00d7 late_fee_per_hr (BR-05)', 'sp_return_vehicle, fn_booking_balance'],
             ['fn_is_vehicle_free', '1 if no BOOKED/ONGOING booking overlaps the period, else 0 (BR-02)', 'sp_create_booking'],
             ['fn_booking_balance', 'Rent + late fee \u2013 successful payments of a completed booking', 'sp_pending_dues (cursor)']],
            [3.6, 9.0, 4.4], size=8.5)
    for k in ('FN_AGE', 'FN_DAYS', 'FN_RENT', 'FN_LATE', 'FN_FREE', 'FN_BALANCE'):
        R.code(fn[k])
    R.figure('funcs1', note='For booking 2 the rent is 3 days \u00d7 2800 = 8400 and the late fee is 4 hours \u00d7 250 = 1000. '
                            'Bookings 12 and 16 still have a balance because one payment is PENDING and one FAILED.')
    R.figure('funcs2')
    R.h2('4.2 Procedures (08_procedures.sql)')
    R.table(['Procedure', 'Requirement', 'Tables changed', 'What it does'],
            [['sp_register_customer', 'FR-01', 'customer', 'Adds a customer with the next sequence ID; duplicates give ORA-20150.'],
             ['sp_create_booking', 'FR-04, FR-07', 'booking, payment', 'Checks branch and overlap, inserts the booking and the advance.'],
             ['sp_pickup_vehicle', 'FR-05', 'booking (vehicle by trigger)', 'BOOKED \u2192 ONGOING and copies the odometer.'],
             ['sp_return_vehicle', 'FR-06, FR-07', 'booking, payment (vehicle by trigger)', 'Records return, computes rent and late fee, inserts FINAL and LATE_FEE.'],
             ['sp_cancel_booking', 'FR-08', 'booking, payment', 'BOOKED \u2192 CANCELLED and inserts the refund by BR-06.']],
            [3.6, 2.3, 4.0, 7.1], size=8.5)
    R.h3('sp_register_customer')
    R.code(sp['SP_REGISTER'])
    R.figure('register')
    R.h3('sp_create_booking')
    R.code(sp['SP_BOOK'])
    R.figure('book', note='Booking 23 is a one-way rental (pickup at branch 1, return at branch 2). The estimated rent is '
                          '2 days \u00d7 Rs 500 for the Scooter category.')
    R.h3('sp_pickup_vehicle')
    R.code(sp['SP_PICKUP'])
    R.figure('pickup', note='The procedure only updates BOOKING. The vehicle status changes to RENTED because of '
                            'trigger trg_booking_status_sync (Section 5.2).')
    R.h3('sp_return_vehicle')
    R.code(sp['SP_RETURN'])
    R.figure('return', note='The return time is passed as a parameter so that a late return can be shown on the same day. '
                            'It is 3 hours late, so the late fee is 3 \u00d7 Rs 50 = Rs 150. The trigger moved vehicle 12 to '
                            'branch 2 (BR-07) and updated its odometer to 9120.')
    R.h3('sp_cancel_booking')
    R.code(sp['SP_CANCEL'])
    R.figure('cancel')

    # ---------------- section 5
    R.h1('5', 'Cursors and Triggers')
    R.h2('5.1 Cursors (09_cursors.sql)')
    R.para('**sp_overdue_report** uses an explicit cursor. It OPENs the cursor, FETCHes one overdue rental at a time, '
           'works out how many hours late it is and the fee so far, and stops when %NOTFOUND becomes true. %ROWCOUNT gives '
           'the number of rows processed, and the cursor is then CLOSEd. This answers the Review I problem of vehicles '
           '"quietly going overdue".')
    R.code(cu['CUR_OVERDUE'])
    R.figure('overdue', note='The output depends on the date and time the procedure is run, because the hours late are '
                             'counted up to SYSDATE.')
    R.para('**sp_pending_dues** uses a cursor FOR loop, which opens, fetches and closes the cursor automatically. For '
           'every completed booking it calls fn_booking_balance and prints only those that still owe money.')
    R.code(cu['CUR_DUES'])
    R.figure('dues')
    R.h2('5.2 Triggers (10_triggers.sql)')
    R.table(['Trigger', 'Timing / event', 'Table', 'Rule'],
            [['trg_customer_validate', 'BEFORE INSERT OR UPDATE OF dob, license_expiry', 'customer', 'BR-01 at registration'],
             ['trg_booking_validate', 'BEFORE INSERT', 'booking', 'BR-01 for the rental period, BR-03'],
             ['trg_booking_status_sync', 'AFTER UPDATE OF status', 'booking', 'Vehicle status, BR-07'],
             ['trg_maintenance_sync', 'AFTER INSERT OR UPDATE OF maint_status', 'maintenance', 'FR-09, BR-03']],
            [4.2, 6.0, 2.6, 4.2], size=8.5)
    R.para('All four are row-level triggers. None of them reads or changes the table it is defined on; they only touch '
           'CUSTOMER and VEHICLE. This avoids the mutating-table error (ORA-04091).', size=10)
    for k in ('TRG_CUSTOMER', 'TRG_BOOKING', 'TRG_SYNC', 'TRG_MAINT'):
        R.code(tg[k])
    R.h2('5.3 Trigger tests')
    R.table(['Action', 'Trigger fired', 'Effect in the database', 'Figure'],
            [['Register an under-age customer / expired licence', 'trg_customer_validate', 'INSERT rejected (ORA-20201, ORA-20202)', fig('trg_customer')],
             ['Book with a licence that ends too early', 'trg_booking_validate', 'INSERT rejected (ORA-20212)', fig('trg_booking')],
             ['Book a vehicle that is RENTED', 'trg_booking_validate', 'INSERT rejected (ORA-20213)', fig('trg_booking')],
             ['sp_pickup_vehicle(23)', 'trg_booking_status_sync', 'Vehicle 12: AVAILABLE \u2192 RENTED', fig('pickup')],
             ['sp_return_vehicle(23, ...)', 'trg_booking_status_sync', 'Vehicle 12: RENTED \u2192 AVAILABLE, branch 1 \u2192 2, odometer 9120', fig('return')],
             ['Maintenance job 11 set to DONE', 'trg_maintenance_sync', 'Vehicle 8: MAINTENANCE \u2192 AVAILABLE', fig('trg_maint')],
             ['New IN_PROGRESS job for vehicle 3', 'trg_maintenance_sync', 'Vehicle 3: AVAILABLE \u2192 MAINTENANCE; booking it is refused', fig('trg_maint')],
             ['IN_PROGRESS job for rented vehicle 2', 'trg_maintenance_sync', 'INSERT rejected (ORA-20221)', fig('trg_maint')]],
            [5.0, 3.9, 6.2, 1.9], size=8.5)
    R.figure('trg_customer')
    R.figure('trg_booking')
    R.figure('trg_maint')

    # ---------------- section 6
    R.h1('6', 'Real-World Transaction Flow')
    R.para('Flow followed by every operation: **User operation \u2192 SQL / PL/SQL \u2192 database processing \u2192 '
           'update or retrieval \u2192 result**.')
    R.table(['User operation', 'Component executed', 'Tables', 'How the database changes', 'Result', 'Fig.'],
            [['Agent registers a customer', 'sp_register_customer, trg_customer_validate', 'customer', 'New row after age/licence check', 'Customer ID', BY_KEY['register']['no']],
             ['Agent books a vehicle', 'sp_create_booking, fn_is_vehicle_free, fn_calculate_rent, trg_booking_validate', 'booking, payment', 'Booking BOOKED + ADVANCE payment', 'Booking ID, rent', BY_KEY['book']['no']],
             ['Customer collects the vehicle', 'sp_pickup_vehicle, trg_booking_status_sync', 'booking, vehicle', 'Booking ONGOING, vehicle RENTED', 'Start odometer', BY_KEY['pickup']['no']],
             ['Customer returns the vehicle', 'sp_return_vehicle, fn_calculate_rent, fn_late_fee, trg_booking_status_sync', 'booking, payment, vehicle', 'COMPLETED, FINAL and LATE_FEE rows, vehicle moved', 'Bill', BY_KEY['return']['no']],
             ['Customer cancels', 'sp_cancel_booking', 'booking, payment', 'CANCELLED + REFUND row', 'Refund amount', BY_KEY['cancel']['no']],
             ['Mechanic starts / closes a job', 'DML on maintenance, trg_maintenance_sync', 'maintenance, vehicle', 'Vehicle MAINTENANCE / AVAILABLE', 'Vehicle status', BY_KEY['trg_maint']['no']],
             ['Manager checks overdue rentals and dues', 'sp_overdue_report, sp_pending_dues', 'booking, customer, vehicle, payment', 'Read only', 'Report lines', f"{BY_KEY['overdue']['no']}, {BY_KEY['dues']['no']}"],
             ['Manager views reports', 'Q1\u2013Q16', 'all tables', 'Read only', 'Result rows', f"{BY_KEY['q1']['no']}\u2013{BY_KEY['q16']['no']}"]],
            [3.2, 4.3, 2.6, 3.2, 2.1, 1.6], size=8)
    R.h2('6.1 Where each business rule of Review I is enforced')
    R.table(['Rule', 'Enforced by', 'Error when broken', 'Evidence'],
            [['BR-01 age \u2265 21, licence valid', 'trg_customer_validate, trg_booking_validate', 'ORA-20201, -20202, -20211, -20212', f"{fig('trg_customer')}, {BY_KEY['trg_booking']['no']}"],
             ['BR-02 no overlapping bookings', 'fn_is_vehicle_free inside sp_create_booking', 'ORA-20103', f"{fig('funcs2')}, {BY_KEY['rules']['no']}"],
             ['BR-03 only AVAILABLE vehicles booked', 'trg_booking_validate', 'ORA-20213', f"{fig('trg_booking')}, {BY_KEY['trg_maint']['no']}"],
             ['BR-04 rent = days \u00d7 rate, min 1 day', 'fn_rental_days, fn_calculate_rent', '\u2013', fig('funcs1')],
             ['BR-05 late fee per started hour', 'fn_late_fee, sp_return_vehicle', '\u2013', f"{fig('q10')}, {BY_KEY['return']['no']}"],
             ['BR-06 refund 100% or 50%', 'sp_cancel_booking', 'ORA-20140 if not BOOKED', f"{fig('q12')}, {BY_KEY['cancel']['no']}"],
             ['BR-07 one-way return moves vehicle', 'trg_booking_status_sync', '\u2013', fig('return')],
             ['BR-08 end odometer \u2265 start', 'ck_booking_odometer, sp_return_vehicle', 'ORA-02290, ORA-20131', fig('ct2')]],
            [4.6, 5.0, 3.9, 3.5], size=8.5)
    R.figure('rules')
    R.h2('6.2 Implementation decisions')
    R.bullets(['**Overlap check in a procedure, not a trigger.** Checking BR-02 needs a SELECT on BOOKING. Doing that in a '
               'row-level trigger on BOOKING can raise the mutating-table error, so it is done in sp_create_booking with fn_is_vehicle_free.',
               '**Part of an hour or day is charged in full.** Review I says "each hour" and "minimum one day"; we charge every started '
               'hour and every started day. Minutes are rounded first so date arithmetic cannot add a false extra unit.',
               '**Advance amount is entered by the agent.** Review I does not fix a percentage, so the advance is a parameter of sp_create_booking.',
               '**A vehicle must be at its pickup branch.** This follows from BR-07: after a one-way rental the vehicle belongs to the return branch.',
               '**Sample data loaded before the triggers.** The history (completed and cancelled rentals) is inserted directly, so the triggers '
               'only act on new transactions.'], size=10)

    # ---------------- section 7
    R.h1('7', 'Viva Preparation: SQL / PL/SQL Understanding')
    R.table(['Likely question', 'Short answer'],
            [['Why are age and licence checked by a trigger and not a CHECK constraint?',
              'They depend on today\'s date. Oracle does not allow SYSDATE in a CHECK constraint, so trg_customer_validate does it before every insert.'],
             ['Procedure vs function in this project?', 'Functions return one value and can be used inside a SELECT (fn_calculate_rent in Figure ' + str(BY_KEY['funcs1']['no']) + '). '
              'Procedures do a full operation and change data (sp_return_vehicle updates BOOKING and inserts PAYMENT rows).'],
             ['Why use sequences?', 'They give unique IDs safely when many agents book at the same time. MAX(id)+1 could give the same ID to two sessions.'],
             ['Which cursors did you use?', 'An explicit cursor (OPEN, FETCH, CLOSE, %NOTFOUND, %ROWCOUNT) in sp_overdue_report and a cursor FOR loop in sp_pending_dues.'],
             ['BEFORE vs AFTER trigger?', 'BEFORE triggers check data before it is written (validation triggers). AFTER triggers react to a change that '
              'has been made, e.g. trg_booking_status_sync updates VEHICLE after the booking status changes.'],
             ['What are :NEW and :OLD?', 'The new and old values of the row being changed. trg_booking_status_sync uses :OLD.status = \'ONGOING\' and '
              ':NEW.status = \'COMPLETED\' to detect a return.'],
             ['What is the mutating-table problem?', 'A row-level trigger cannot query or change the table that fired it (ORA-04091). That is why '
              'the overlap check is inside sp_create_booking.'],
             ['What does ON DELETE CASCADE do here?', 'Deleting a booking deletes its payments automatically (Figure ' + str(BY_KEY['dml_delete']['no']) +
              '). Other foreign keys use the default, so a customer with bookings cannot be deleted (ORA-02292).'],
             ['WHERE vs HAVING?', 'WHERE filters rows before grouping (role = \'AGENT\' in Q15); HAVING filters groups after aggregation (COUNT >= 2).'],
             ['Why LEFT JOIN in Q2 and Q8?', 'To keep branches or categories with no matching rows; an inner join would drop the Hitech City branch.'],
             ['Why is there a COMMIT in each procedure?', 'Each procedure is one business transaction. If any step fails, RAISE_APPLICATION_ERROR '
              'stops it before COMMIT, so a booking is never saved without its advance payment.']],
            [5.4, 11.6], size=8.5)

    # ---------------- section 8
    R.h1('8', 'Conclusion')
    R.para('The Review I design has been implemented in Oracle without changes to the schema. The 8 tables and their 85 '
           'constraints protect the data at the lowest level. Six functions compute the derived values (age, rental days, '
           'rent, late fee, availability and balance). Five procedures carry out the main business operations, and two '
           'cursor procedures produce the overdue and pending-dues reports. Four triggers keep vehicle status, branch and '
           'odometer correct without the application having to remember it. The 16 queries cover the reports listed in '
           'Review I. For Review III we will connect a Python application to this database so that agents can perform '
           'these transactions from a simple interface.')
    R.save(out_path)


def build_checklist(out_path):
    R = Report()
    d = R.doc
    p = d.add_paragraph()
    para_fmt(p, after=2)
    add_run(p, 'Review II \u2013 Screenshot Checklist', bold=True, size=18, color=GREEN)
    R.para('Work through this list from top to bottom in one SQL*Plus session (Run SQL Command Line). The order is the '
           'execution order, which is slightly different from the order of the figures in the report.')
    R.h2('Before you start')
    R.bullets([
        'Start from an empty schema. If you ran the scripts before, run `@99_cleanup.sql` first (ignore "does not exist" errors).',
        'Put all files of the sql/ folder in one folder and start SQL*Plus from that folder, or use the full path, e.g. `@C:\\dbms\\sql\\01_ddl.sql`.',
        'Make the console window wide (about 150 columns) so rows do not wrap: Properties \u2192 Layout \u2192 Window size width 150.',
        'Run `@00_setup_display.sql` at the start of every new session (after reconnecting as well).',
        'For each figure: clear the screen (`CLEAR SCREEN`), paste the block from the file, press Enter, then take the screenshot '
        '(Windows: Win + Shift + S). Include the typed command and the full output; crop away the empty part of the window.',
        'Save as screenshots/fig_NN.png using the figure number below (e.g. fig_14.png). Then either paste each picture into the '
        'placeholder box in the DOCX, or run `python3 tools/build_report.py` to insert them automatically.',
        'If any output differs from "What you should see", do not edit the screenshot. Note the difference and fix the script or ask for help.'])
    order = sorted([f for f in FIGURES if 'step' in f], key=lambda f: (f['step'], f['no'] if f['step'] != 3 else 0))
    # step 3 figures follow file order, others follow figure numbers
    step3 = ['sc_desc', 'sc_count', 'sc_fk', 'data_counts', 'data_vehicle', 'data_booking']
    order = [f for f in order if f['step'] < 3] + [BY_KEY[k] for k in step3] + [f for f in order if f['step'] > 3]
    steps = {1: '@00_setup_display.sql then @01_ddl.sql',
             2: '@02_insert_data.sql (no screenshot; the last lines must be "3 rows updated." / "Commit complete.")',
             3: '03_check_schema_data.sql \u2013 run each block separately',
             4: '04_constraint_tests.sql \u2013 two blocks, every statement must fail',
             5: '05_queries.sql \u2013 run Q1 to Q16 one at a time',
             6: '06_dml_demo.sql \u2013 three blocks',
             7: '@07_functions.sql, @08_procedures.sql, @09_cursors.sql, @10_triggers.sql (no screenshot; each object must say "created.")',
             8: '11_transactions.sql \u2013 run the blocks strictly in this order (IDs come from sequences)'}
    rows, last = [], None
    n = 0
    for f in order:
        if f['step'] != last:
            for s in range(last + 1 if last else 1, f['step'] + 1):
                if s in (2, 7) or s == f['step']:
                    rows.append(('step', s))
            last = f['step']
        rows.append(('fig', f))
    for kind, val in rows:
        if kind == 'step':
            R.h2(f'Step {val}: {steps[val]}')
            continue
        f = val
        n += 1
        src = f.get('src')
        R.h3(f"{n}.  Figure {f['no']} \u2013 {f['caption']}",
             sub=f"Save as **screenshots/fig_{f['no']:02d}.png**" + (f"  |  Source: {src[0]}, block {src[1]}" if src else ''))
        cmd = chunks(src[0])[src[1]] if src else f['run']
        R.code(cmd, size=7)
        R.para('**What you should see:** ' + f['expect'], size=9.5)
    R.save_checklist = True
    d.save(out_path)


if __name__ == '__main__':
    out = os.path.join(ROOT, 'Vehicle_Rental_Review2_Report.docx')
    build_report(out)
    build_checklist(os.path.join(ROOT, 'Review2_Screenshot_Checklist.docx'))
    print('written', out)
